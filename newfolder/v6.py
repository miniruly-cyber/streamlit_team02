# Requirements (install first):
#   pip install streamlit python-docx reportlab langchain langchain-openai langchain-google-genai python-dotenv googletrans

import os, io, json, time, textwrap, re, datetime, urllib.parse
import streamlit as st
from typing import Optional, Tuple, List, Dict
import base64

# ===== 문서 생성을 위한 라이브러리 =====
try:
    from docx import Document
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    DOC_LIBS_AVAILABLE = True
except ImportError:
    DOC_LIBS_AVAILABLE = False

# ===== LangChain imports (조건부) =====
try:
    from langchain_openai import ChatOpenAI
    from langchain_google_genai import ChatGoogleGenerativeAI
    from langchain.prompts import ChatPromptTemplate
    from langchain.memory import ConversationBufferMemory
    from langchain.chains import LLMChain
    LANGCHAIN_AVAILABLE = True
except ImportError:
    LANGCHAIN_AVAILABLE = False

# ===== 번역을 위한 라이브러리 =====
try:
    from googletrans import Translator
    TRANSLATOR_AVAILABLE = True
except ImportError:
    TRANSLATOR_AVAILABLE = False

# ================= 기본 설정 =================
st.set_page_config(
    page_title="AI 자기소개서 코치",
    page_icon="🎯",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# 개선된 스타일
st.markdown("""
<style>
.main .block-container {
    max-width: 900px;
    padding: 2rem 1rem;
}

.header-container {
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    border-radius: 20px;
    padding: 2rem;
    margin-bottom: 2rem;
    text-align: center;
    color: white;
    box-shadow: 0 10px 30px rgba(0,0,0,0.1);
}

.header-title {
    font-size: 2.5rem;
    font-weight: 800;
    margin-bottom: 0.5rem;
}

.header-subtitle {
    font-size: 1.1rem;
    opacity: 0.9;
    font-weight: 300;
}

.chat-container {
    background: white;
    border-radius: 20px;
    padding: 1.5rem;
    margin-bottom: 1rem;
    box-shadow: 0 5px 20px rgba(0,0,0,0.05);
    border: 1px solid #f0f0f0;
    height: 500px;
    overflow-y: auto;
}

.message-bubble {
    margin: 1rem 0;
    display: flex;
    align-items: flex-start;
}

.message-bubble.user {
    justify-content: flex-end;
}

.message-bubble.bot {
    justify-content: flex-start;
}

.message-content {
    max-width: 75%;
    padding: 1rem 1.5rem;
    border-radius: 20px;
    font-size: 0.95rem;
    line-height: 1.6;
    word-wrap: break-word;
}

.message-bubble.user .message-content {
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    color: white;
    border-bottom-right-radius: 5px;
}

.message-bubble.bot .message-content {
    background: #f8fafc;
    color: #1a202c;
    border: 1px solid #e2e8f0;
    border-bottom-left-radius: 5px;
}

.stTabs [data-baseweb="tab-list"] {
    gap: 0.5rem;
    background: #f8fafc;
    padding: 0.5rem;
    border-radius: 15px;
}

.stTabs [data-baseweb="tab"] {
    height: 3rem;
    padding: 0 1.5rem;
    background: white;
    border-radius: 10px;
    border: none;
    font-weight: 500;
}

.stTabs [aria-selected="true"] {
    background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
    color: white;
}

.feature-card {
    background: white;
    border-radius: 15px;
    padding: 1.5rem;
    margin: 1rem 0;
    box-shadow: 0 5px 15px rgba(0,0,0,0.05);
    border: 1px solid #e2e8f0;
}

.guideline-card {
    background: linear-gradient(135deg, #e0f2f1 0%, #f3e5f5 100%);
    border-radius: 15px;
    padding: 1.5rem;
    margin: 1rem 0;
    border-left: 4px solid #667eea;
}

.guideline-item {
    background: white;
    border-radius: 10px;
    padding: 1rem;
    margin: 0.5rem 0;
    box-shadow: 0 2px 8px rgba(0,0,0,0.05);
}

.file-item {
    background: white;
    border-radius: 10px;
    padding: 1rem;
    margin: 0.5rem 0;
    box-shadow: 0 2px 8px rgba(0,0,0,0.05);
    border: 1px solid #e2e8f0;
}
</style>
""", unsafe_allow_html=True)

# ================= 상태 초기화 =================
if "initialized" not in st.session_state:
    st.session_state.initialized = True
    st.session_state.msgs = []
    st.session_state.settings = {
        "provider": "openai",
        "model": "gpt-4o-mini",
        "tone": "정중하고 간결한",
        "length": 800,
        "temperature": 0.7,
        "openai_key": os.getenv("OPENAI_API_KEY", ""),
        "gemini_key": os.getenv("GEMINI_API_KEY", ""),
        "save_dir": "./AI_CoverLetter_Storage",
        "font_family": "Nanum Gothic",
        "enable_translation": False,
        "use_free_model": True
    }
    
    # 저장 디렉토리 생성
    os.makedirs(st.session_state.settings["save_dir"], exist_ok=True)
    
    # 초기 메시지 추가
    st.session_state.msgs.append({
        "role": "bot",
        "content": "안녕하세요! AI 자기소개서 코치입니다. 🎯\n\n어떤 도움이 필요하신가요?",
        "timestamp": datetime.datetime.now().strftime("%H:%M")
    })

if "saved_files" not in st.session_state:
    st.session_state.saved_files = []

if LANGCHAIN_AVAILABLE and "lc_memory" not in st.session_state:
    st.session_state.lc_memory = ConversationBufferMemory(memory_key="chat_history", return_messages=True)

# ================= 유틸리티 함수 =================
def now_hhmm():
    return datetime.datetime.now().strftime("%H:%M")

def timestamp():
    return datetime.datetime.now().strftime("%Y%m%d_%H%M%S")

def translate_to_english(text: str) -> str:
    """텍스트를 영어로 번역"""
    if not TRANSLATOR_AVAILABLE:
        return "번역 기능을 사용하려면 googletrans 라이브러리를 설치해주세요."
    
    try:
        translator = Translator()
        result = translator.translate(text, src='ko', dest='en')
        return result.text
    except Exception as e:
        return f"번역 중 오류가 발생했습니다: {str(e)}"

def get_free_ai_response(user_message: str) -> str:
    """무료 AI 응답 생성"""
    response_templates = {
        "마케팅": """📊 **마케팅 직무 자기소개서 작성 가이드**

**1. 핵심 역량 강조**
- 데이터 분석 및 인사이트 도출 능력
- 창의적 캠페인 기획 경험  
- 디지털 마케팅 도구 활용 능력

**2. 구체적 성과 제시**
- "매출 20% 증가" 같은 정량적 결과
- "CTR 3% 향상" 등 구체적 지표
- "신규 고객 1,000명 확보" 등 명확한 수치

**3. 경험 서술 방법**
- STAR 기법 활용 (상황-과제-행동-결과)
- 문제 해결 과정과 결과 중심
- 팀워크와 리더십 경험 포함""",
        
        "개발": """💻 **개발 직무 자기소개서 작성 가이드**

**1. 기술 스택 명시**
- 사용 가능한 프로그래밍 언어
- 프레임워크 및 라이브러리 경험
- 데이터베이스 및 클라우드 경험

**2. 프로젝트 경험 상세화**
- 개발한 서비스의 규모와 성과
- 해결한 기술적 문제와 방법
- 코드 품질 향상을 위한 노력

**3. 성장 의지 표현**
- 지속적 학습과 기술 트렌드 관심
- 오픈소스 기여나 개인 프로젝트
- 새로운 기술에 대한 도전 의지""",
        
        "영업": """🎯 **영업 직무 자기소개서 작성 가이드**

**1. 영업 성과 강조**
- 목표 달성률과 매출 기여도
- 신규 고객 개발 성과
- 기존 고객과의 관계 유지 성과

**2. 커뮤니케이션 능력**
- 고객 니즈 파악 및 솔루션 제안
- 설득력 있는 프레젠테이션 경험
- 다양한 이해관계자와의 협업

**3. 시장 이해도**
- 업계 트렌드 및 경쟁사 분석
- 고객사 비즈니스 모델 이해
- 시장 변화에 대한 대응 능력"""
    }
    
    user_lower = user_message.lower()
    
    if "마케팅" in user_lower:
        return response_templates["마케팅"]
    elif any(word in user_lower for word in ["개발", "프로그래밍", "코딩", "IT"]):
        return response_templates["개발"]
    elif "영업" in user_lower:
        return response_templates["영업"]
    elif any(word in user_lower for word in ["첨삭", "피드백", "검토"]):
        return """✏️ **자기소개서 첨삭 포인트**

**1. 구조와 논리성**
- 도입-본론-결론의 명확한 구성
- 각 문단 간의 논리적 연결
- 핵심 메시지의 일관성

**2. 내용의 구체성**  
- 추상적 표현을 구체적 사례로 변경
- 성과와 결과를 수치로 표현
- 개인의 독특한 경험과 강점 부각

**3. 문장과 표현**
- 간결하고 명확한 문장 구조
- 반복되는 표현 제거
- 전문적이면서도 자연스러운 어조

📎 파일을 업로드하시면 더 구체적인 첨삭을 도와드릴 수 있습니다!"""
    else:
        return """🎯 **자기소개서 작성을 도와드릴게요!**

**효과적인 질문 예시:**
- "마케팅 직무 자기소개서 작성법 알려주세요"
- "IT 개발자 자소서에서 강조해야 할 점은?"
- "영업 직무 경험을 어떻게 표현하면 좋을까요?"
- "제 자기소개서 첨삭해주세요" (파일 첨부)

**기본 작성 원칙:**
1. **STAR 기법** - 상황, 과제, 행동, 결과
2. **구체적 수치** - 성과를 정량적으로 표현
3. **차별화** - 나만의 독특한 경험과 강점"""

def get_ai_response(user_message: str, uploaded_file=None) -> str:
    """AI 응답 생성"""
    settings = st.session_state.settings
    
    # 무료 모델 사용 또는 API 키 없음
    if settings["use_free_model"] or (not settings["openai_key"] and not settings["gemini_key"]):
        if uploaded_file is not None:
            try:
                file_content = uploaded_file.read().decode('utf-8')
                return f"""📋 **업로드된 자기소개서 첨삭**

**원본 내용 (일부):**
{file_content[:200]}...

**첨삭 의견:**
{get_free_ai_response("첨삭")}

💡 더 정교한 첨삭을 위해서는 설정에서 API 키를 설정해주세요."""
            except Exception as e:
                return f"파일 읽기 오류: {str(e)}"
        else:
            return get_free_ai_response(user_message)
    
    if not LANGCHAIN_AVAILABLE:
        return get_free_ai_response(user_message)
    
    try:
        # 모델 선택
        if settings["provider"] == "openai" and settings["openai_key"]:
            llm = ChatOpenAI(
                api_key=settings["openai_key"],
                model=settings["model"],
                temperature=settings["temperature"]
            )
        elif settings["provider"] == "gemini" and settings["gemini_key"]:
            llm = ChatGoogleGenerativeAI(
                google_api_key=settings["gemini_key"],
                model="gemini-pro",
                temperature=settings["temperature"]
            )
        else:
            return get_free_ai_response(user_message)
        
        # 파일 첨삭 모드
        if uploaded_file is not None:
            file_content = uploaded_file.read().decode('utf-8')
            prompt_text = f"""
            다음 자기소개서를 전문가 관점에서 첨삭해주세요:
            
            [자기소개서 내용]
            {file_content}
            
            [사용자 질문]
            {user_message}
            
            다음 관점에서 상세한 피드백을 제공해주세요:
            1. 구조와 논리성
            2. 내용의 구체성과 차별화  
            3. 문장 표현과 어법
            4. 개선 제안사항
            """
        else:
            prompt_text = user_message
        
        # 프롬프트 설정
        system_prompt = f"""
        당신은 자기소개서 작성 전문 코치입니다.
        - 톤: {settings["tone"]}
        - 목표 길이: 약 {settings["length"]}자
        - 구체적이고 실용적인 조언을 제공하세요
        - STAR 기법을 활용한 경험 서술을 권장하세요
        - 정량적 성과와 구체적 사례를 강조하세요
        """
        
        prompt = ChatPromptTemplate.from_messages([
            ("system", system_prompt),
            ("human", "{input}")
        ])
        
        chain = LLMChain(llm=llm, prompt=prompt, memory=st.session_state.lc_memory)
        response = chain.run(input=prompt_text)
        
        # 영문 변환 기능
        if settings["enable_translation"] and not uploaded_file:
            english_version = translate_to_english(response)
            response += f"\n\n---\n**영문 버전:**\n{english_version}"
        
        return response
        
    except Exception as e:
        return f"오류가 발생했습니다: {str(e)}\n\n{get_free_ai_response(user_message)}"

# ================= 문서 생성 함수들 =================
def create_txt(content: str, filename: str) -> str:
    """TXT 파일 생성"""
    try:
        filepath = os.path.join(st.session_state.settings["save_dir"], f"{filename}.txt")
        with open(filepath, 'w', encoding='utf-8') as f:
            f.write(f"자기소개서\n{'='*20}\n\n")
            f.write(content)
        return filepath
    except Exception as e:
        st.error(f"TXT 생성 중 오류: {str(e)}")
        return None

def create_docx(content: str, filename: str) -> str:
    """DOCX 파일 생성"""
    if not DOC_LIBS_AVAILABLE:
        return None
    
    try:
        filepath = os.path.join(st.session_state.settings["save_dir"], f"{filename}.docx")
        doc = Document()
        
        # 제목 추가
        title = doc.add_heading('자기소개서', 0)
        title.alignment = 1
        
        # 본문 추가
        for line in content.split('\n'):
            if line.strip():
                doc.add_paragraph(line)
        
        doc.save(filepath)
        return filepath
    except Exception as e:
        st.error(f"DOCX 생성 중 오류: {str(e)}")
        return None

def create_pdf(content: str, filename: str) -> str:
    """PDF 파일 생성"""
    if not DOC_LIBS_AVAILABLE:
        return None
    
    try:
        filepath = os.path.join(st.session_state.settings["save_dir"], f"{filename}.pdf")
        doc = SimpleDocTemplate(filepath, pagesize=letter)
        
        styles = getSampleStyleSheet()
        story = []
        
        # 제목
        title_style = ParagraphStyle(
            'Title',
            parent=styles['Heading1'],
            fontSize=18,
            spaceAfter=30,
            alignment=1
        )
        
        story.append(Paragraph("자기소개서", title_style))
        story.append(Spacer(1, 12))
        
        # 본문
        for line in content.split('\n'):
            if line.strip():
                story.append(Paragraph(line, styles['Normal']))
            else:
                story.append(Spacer(1, 6))
        
        doc.build(story)
        return filepath
    except Exception as e:
        st.error(f"PDF 생성 중 오류: {str(e)}")
        return None

def save_conversation(file_format: str, filename: str) -> str:
    """대화 내용을 파일로 저장"""
    conversation_text = ""
    for msg in st.session_state.msgs:
        if msg["role"] == "user":
            conversation_text += f"👤 사용자: {msg['content']}\n\n"
        else:
            conversation_text += f"🤖 AI 코치: {msg['content']}\n\n"
        conversation_text += "---\n\n"
    
    # 파일 형식에 따라 저장
    if file_format == "pdf":
        filepath = create_pdf(conversation_text, filename)
    elif file_format == "docx":
        filepath = create_docx(conversation_text, filename)
    else:
        filepath = create_txt(conversation_text, filename)
    
    if filepath and os.path.exists(filepath):
        file_info = {
            "name": os.path.basename(filepath),
            "path": filepath,
            "created": datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
            "size": os.path.getsize(filepath)
        }
        if file_info not in st.session_state.saved_files:
            st.session_state.saved_files.append(file_info)
        return filepath
    return None

def get_saved_files() -> List[Dict]:
    """저장된 파일 목록 반환"""
    saved_files = []
    save_dir = st.session_state.settings["save_dir"]
    
    if os.path.exists(save_dir):
        for filename in os.listdir(save_dir):
            filepath = os.path.join(save_dir, filename)
            if os.path.isfile(filepath):
                file_info = {
                    "name": filename,
                    "path": filepath,
                    "created": datetime.datetime.fromtimestamp(
                        os.path.getctime(filepath)
                    ).strftime("%Y-%m-%d %H:%M:%S"),
                    "size": os.path.getsize(filepath)
                }
                saved_files.append(file_info)
    
    saved_files.sort(key=lambda x: x["created"], reverse=True)
    return saved_files

# ================= UI 렌더링 함수들 =================
def render_header():
    """헤더 렌더링"""
    st.markdown("""
    <div class="header-container">
        <div class="header-title">🎯 AI 자기소개서 코치</div>
        <div class="header-subtitle">전문 AI가 도와드리는 맞춤형 자기소개서 작성 서비스</div>
    </div>
    """, unsafe_allow_html=True)

def render_guidelines():
    """질문 가이드라인 렌더링"""
    st.markdown("""
    <div class="guideline-card">
        <h3>💡 효과적인 질문 방법 가이드</h3>
        
        <div class="guideline-item">
            <strong>🎯 직무별 맞춤 질문</strong>
            <p>• "마케팅 직무 자기소개서 작성법 알려주세요"<br>
            • "IT 개발자로 지원할 때 강조해야 할 점은?"<br>
            • "영업 직무 경험을 어떻게 어필하면 좋을까요?"</p>
        </div>
        
        <div class="guideline-item">
            <strong>📝 상황별 구체적 질문</strong>
            <p>• "신입사원 자소서에서 학교 프로젝트 경험 어떻게 쓸까요?"<br>
            • "경력직 이직 시 이직 사유 어떻게 표현하면 좋을까요?"<br>
            • "다른 분야에서 전환할 때 어떻게 어필해야 하나요?"</p>
        </div>
        
        <div class="guideline-item">
            <strong>✏️ 작성 기법 문의</strong>
            <p>• "STAR 기법으로 경험을 어떻게 구조화하나요?"<br>
            • "성과를 수치로 표현하는 방법 알려주세요"<br>
            • "자소서 길이는 어느 정도가 적당한가요?"</p>
        </div>
        
        <div class="guideline-item">
            <strong>🔍 첨삭 및 피드백 요청</strong>
            <p>• "제 자기소개서 첨삭해주세요" + 파일 첨부<br>
            • "이 표현이 자연스러운지 확인해주세요"<br>
            • "더 임팩트 있게 표현하는 방법은?"</p>
        </div>
    </div>
    """, unsafe_allow_html=True)

def render_chat_tab():
    """채팅 탭 렌더링"""
    render_header()
    
    # 가이드라인 표시
    with st.expander("💡 질문 가이드라인 보기", expanded=False):
        render_guidelines()
    
    # 채팅 메시지 표시
    st.markdown('<div class="chat-container">', unsafe_allow_html=True)
    
    # 메시지 표시
    for msg in st.session_state.msgs:
        role_class = "user" if msg["role"] == "user" else "bot"
        
        st.markdown(f"""
        <div class="message-bubble {role_class}">
            <div class="message-content">
                {msg["content"].replace(chr(10), '<br>')}
            </div>
        </div>
        """, unsafe_allow_html=True)
    
    st.markdown('</div>', unsafe_allow_html=True)
    
    # 파일 업로드
    uploaded_file = st.file_uploader(
        "📎 자기소개서 파일 첨부 (첨삭용)",
        type=['txt', 'docx'],
        help="TXT 또는 DOCX 파일을 업로드하면 첨삭을 도와드립니다."
    )
    
    # 입력 및 전송
    with st.form(key="chat_form", clear_on_submit=True):
        col1, col2 = st.columns([5, 1])
        
        with col1:
            user_input = st.text_input(
                "메시지를 입력하세요...",
                placeholder="예: 마케팅 직무 자기소개서 작성법을 알려주세요",
                label_visibility="collapsed"
            )
        
        with col2:
            submit = st.form_submit_button("전송", use_container_width=True, type="primary")
        
        if submit and user_input:
            # 사용자 메시지 추가
            st.session_state.msgs.append({
                "role": "user", 
                "content": user_input,
                "timestamp": now_hhmm()
            })
            
            # AI 응답 생성
            with st.spinner("AI가 답변을 생성중입니다..."):
                ai_response = get_ai_response(user_input, uploaded_file)
                st.session_state.msgs.append({
                    "role": "bot",
                    "content": ai_response,
                    "timestamp": now_hhmm()
                })
            
            st.rerun()

def render_settings_tab():
    """설정 탭 렌더링"""
    st.markdown("""
    <div class="feature-card">
        <h2>⚙️ AI 모델 및 응답 설정</h2>
    </div>
    """, unsafe_allow_html=True)
    
    settings = st.session_state.settings
    
    # 무료 모드 선택
    st.markdown("### 🆓 모델 사용 방식")
    use_free = st.checkbox(
        "무료 모드 사용 (API 키 없이 기본 가이드 제공)",
        value=settings["use_free_model"]
    )
    
    if not use_free:
        col1, col2 = st.columns(2)
        
        with col1:
            st.markdown("### 🔑 API 키 설정")
            
            provider = st.selectbox(
                "AI 제공업체",
                ["openai", "gemini"],
                index=0 if settings["provider"] == "openai" else 1
            )
            
            if provider == "openai":
                openai_key = st.text_input(
                    "OpenAI API Key",
                    value=settings["openai_key"],
                    type="password",
                    help="OpenAI API 키를 입력하세요"
                )
                
                model = st.selectbox(
                    "OpenAI 모델",
                    ["gpt-4o-mini", "gpt-4o", "gpt-3.5-turbo"],
                    index=["gpt-4o-mini", "gpt-4o", "gpt-3.5-turbo"].index(settings["model"])
                )
            else:
                gemini_key = st.text_input(
                    "Google Gemini API Key",
                    value=settings["gemini_key"],
                    type="password",
                    help="Google Gemini API 키를 입력하세요"
                )
                model = "gemini-pro"
                openai_key = settings["openai_key"]
